"""
Independent artefact verification.

The test suite checks that the *code* behaves. This checks that the *artefacts
in the repository are what the reports claim they are* - a different question,
and the one a reader who did not run the pipeline actually cares about.

It re-derives things rather than trusting them:

* dataset file SHA-256 against the recorded provenance
* the three splits are pairwise disjoint at feature-vector level (leakage)
* the persisted model's feature list contains no leakage column
* **the headline test metrics, recomputed from the saved model and the data**,
  agreed to 1e-6 against what the reports say
* every figure referenced by a report exists and is a readable PNG
* every notebook is fully executed with no error outputs
* both decks have the slide count their content file declares
* local links in README.md resolve

Run with ``python scripts/verify_project.py``. Exit code 0 means every check
passed; any failure is printed with its detail and the exit code is 1. Results
are written to ``reports/metrics/artifact_verification.json``.
"""

from __future__ import annotations

import hashlib
import json
import re
import struct
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

from src import config  # noqa: E402

RESULTS: list[dict[str, object]] = []


def check(name: str):
    """Run one check, record pass/fail, never let an exception stop the rest."""
    def decorator(fn):
        try:
            detail = fn()
            RESULTS.append({"check": name, "status": "pass", "detail": detail})
            print(f"  PASS  {name}" + (f" - {detail}" if detail else ""))
        except AssertionError as exc:
            RESULTS.append({"check": name, "status": "FAIL", "detail": str(exc)})
            print(f"  FAIL  {name} - {exc}")
        except Exception as exc:  # noqa: BLE001 - report, do not abort
            RESULTS.append({"check": name, "status": "ERROR",
                            "detail": f"{type(exc).__name__}: {exc}"})
            print(f"  ERROR {name} - {type(exc).__name__}: {exc}")
        return fn
    return decorator


def skip(name: str, reason: str) -> None:
    RESULTS.append({"check": name, "status": "skip", "detail": reason})
    print(f"  SKIP  {name} - {reason}")


def png_dimensions(path: Path) -> tuple[int, int]:
    """Read a PNG header directly - a truncated file fails here."""
    data = path.read_bytes()
    assert data[:8] == b"\x89PNG\r\n\x1a\n", f"{path.name} is not a PNG"
    # The IEND chunk is the final 12 bytes: 4-byte length (zero), the literal
    # b"IEND", then a 4-byte CRC - so the type field sits at [-8:-4], not [-12:-8].
    assert data[-8:-4] == b"IEND", f"{path.name} is truncated (no IEND chunk)"
    width, height = struct.unpack(">II", data[16:24])
    return width, height


# --------------------------------------------------------------------------- #
def main() -> int:
    print("Verifying repository artefacts\n")
    metrics = config.METRICS_DIR
    data_present = (config.RAW_DIR / config.TRAIN_FILE).exists()

    # ---------------- repository shape ----------------
    @check("required files present")
    def _required():
        required = ["README.md", "requirements.txt", "requirements-lock.txt",
                    "environment.yml", ".gitignore", ".gitattributes", "LICENSE",
                    "pytest.ini", "app/streamlit_app.py"]
        missing = [name for name in required if not (ROOT / name).is_file()]
        assert not missing, f"missing: {missing}"
        return f"{len(required)} files"

    # ---------------- dataset integrity ----------------
    if data_present:
        @check("dataset SHA-256 matches recorded provenance")
        def _hashes():
            from src import data_loader
            recorded = data_loader.file_checksums()
            assert recorded, "no checksums recorded"
            for filename, expected in recorded.items():
                actual = hashlib.sha256(
                    (config.RAW_DIR / filename).read_bytes()).hexdigest()
                assert actual == expected, f"{filename}: {actual[:12]} != {expected[:12]}"
            return f"{len(recorded)} files"
    else:
        skip("dataset SHA-256 matches recorded provenance", "data/raw/ not present")

    # ---------------- split disjointness ----------------
    if data_present:
        @check("train / validation / test are pairwise disjoint")
        def _disjoint():
            from src import preprocessing
            train, val, test, _ = preprocessing.split_published(verbose=False)
            predictors = [c for c in train.columns
                          if c not in (*config.LEAKAGE_COLUMNS, "partition")]

            def sig(frame):
                return set(map(tuple, frame[predictors].itertuples(index=False, name=None)))

            s_train, s_val, s_test = sig(train), sig(val), sig(test)
            assert not s_train & s_val, f"{len(s_train & s_val)} shared train/val vectors"
            assert not s_train & s_test, f"{len(s_train & s_test)} shared train/test vectors"
            assert not s_val & s_test, f"{len(s_val & s_test)} shared val/test vectors"
            return f"train {len(train):,} / val {len(val):,} / test {len(test):,}"
    else:
        skip("train / validation / test are pairwise disjoint", "data/raw/ not present")

    # ---------------- model schema ----------------
    reference_path = config.MODELS_DIR / "feature_reference.json"

    @check("persisted feature schema excludes every leakage column")
    def _schema():
        assert reference_path.exists(), "models/feature_reference.json is missing"
        reference = json.loads(reference_path.read_text(encoding="utf-8"))
        columns = set(reference.get("numeric", {})) | set(reference.get("categorical", {}))
        assert columns, "feature reference lists no columns"
        leaked = columns & set(config.LEAKAGE_COLUMNS)
        assert not leaked, f"leakage columns in feature schema: {sorted(leaked)}"

        # The example rows are what the app pre-fills; they must not carry the
        # target either, or a screenshot of the app would be showing the answer.
        for name, example in reference.get("examples", {}).items():
            leaked_example = set(example) & set(config.LEAKAGE_COLUMNS)
            assert not leaked_example, f"example '{name}' carries {sorted(leaked_example)}"
        return f"{len(columns)} predictor columns, none leaking"

    # ---------------- independent metric recomputation ----------------
    deployment_path = config.MODELS_DIR / "deployment.json"
    test_metrics_path = metrics / "test_metrics_main.csv"

    if data_present and deployment_path.exists() and test_metrics_path.exists():
        @check("headline test metrics recompute from the saved model")
        def _recompute():
            from sklearn.metrics import (average_precision_score, f1_score,
                                         precision_score, recall_score, roc_auc_score)

            from src import predict, preprocessing
            deployment = json.loads(deployment_path.read_text(encoding="utf-8"))
            model = predict.load_model()
            _, _, test, _ = preprocessing.split_published(verbose=False)
            X_test, y_test = preprocessing.split_xy(test)

            proba = model.pipeline.predict_proba(X_test)[:, 1]
            threshold = float(deployment["threshold"])
            pred = (proba >= threshold).astype(int)
            recomputed = {
                "precision": precision_score(y_test, pred, zero_division=0),
                "recall": recall_score(y_test, pred, zero_division=0),
                "f1": f1_score(y_test, pred, zero_division=0),
                "roc_auc": roc_auc_score(y_test, proba),
                "pr_auc": average_precision_score(y_test, proba),
            }
            reported = pd.read_csv(test_metrics_path, index_col=0).loc[
                deployment["model_key"]]
            for key, value in recomputed.items():
                np.testing.assert_allclose(
                    value, float(reported[key]), rtol=1e-6, atol=1e-7,
                    err_msg=f"{key}: recomputed {value:.8f} vs reported {reported[key]}")
            return (f"F1 {recomputed['f1']:.6f}, PR-AUC {recomputed['pr_auc']:.6f} "
                    f"on n={len(test):,}")
    else:
        skip("headline test metrics recompute from the saved model",
             "needs data/raw/, models/deployment.json and test metrics")

    # ---------------- figures ----------------
    @check("every figure referenced by a report exists and is a valid PNG")
    def _figures():
        referenced: set[str] = set()
        for markdown in list(config.REPORTS_DIR.glob("*.md")) + \
                [ROOT / "README.md"] + list((ROOT / "presentations").glob("*.md")):
            if markdown.exists():
                referenced |= set(re.findall(r"(fig\d{2}[A-Za-z0-9_]*)\.png",
                                             markdown.read_text(encoding="utf-8")))
        assert referenced, "no figures referenced anywhere"
        missing, broken = [], []
        for stem in sorted(referenced):
            path = config.FIGURES_DIR / f"{stem}.png"
            if not path.exists():
                missing.append(stem)
                continue
            try:
                width, height = png_dimensions(path)
                if width < 200 or height < 150:
                    broken.append(f"{stem} ({width}x{height})")
            except AssertionError as exc:
                broken.append(str(exc))
        assert not missing, f"referenced but absent: {missing}"
        assert not broken, f"unreadable or undersized: {broken}"
        return f"{len(referenced)} figures"

    # ---------------- notebooks ----------------
    @check("every notebook is fully executed with no errors")
    def _notebooks():
        notebooks = sorted(config.NOTEBOOKS_DIR.glob("*.ipynb"))
        assert len(notebooks) == 6, f"expected 6 notebooks, found {len(notebooks)}"
        for path in notebooks:
            nb = json.loads(path.read_text(encoding="utf-8"))
            code = [c for c in nb["cells"] if c["cell_type"] == "code"]
            unexecuted = [c for c in code if c.get("execution_count") is None]
            errored = [c for c in code if any(
                o.get("output_type") == "error" for o in c.get("outputs", []))]
            assert not unexecuted, f"{path.name}: {len(unexecuted)} unexecuted cells"
            assert not errored, f"{path.name}: {len(errored)} cells raised"
        return f"{len(notebooks)} notebooks, all executed cleanly"

    # ---------------- presentations ----------------
    @check("both decks match their content files")
    def _decks():
        try:
            from pptx import Presentation
        except ImportError:
            raise AssertionError("python-pptx not installed")
        pairs = [("technical_presentation_content.md",
                  "Arne_Ramos_AI_Capstone_Technical_Presentation.pptx"),
                 ("executive_presentation_content.md",
                  "Arne_Ramos_AI_Capstone_Executive_Presentation.pptx")]
        summary = []
        for source_name, deck_name in pairs:
            source = ROOT / "presentations" / source_name
            deck = ROOT / "deliverables" / deck_name
            assert source.exists(), f"{source_name} missing"
            assert deck.exists(), f"{deck_name} missing"
            declared = len(re.findall(r"^## Slide ", source.read_text(encoding="utf-8"), re.M))
            built = len(Presentation(deck).slides)
            # The built deck adds a title slide ahead of the numbered content.
            assert built in (declared, declared + 1), \
                f"{deck_name}: {built} slides vs {declared} declared"
            summary.append(f"{deck_name.split('_')[-2]} {built}")
        return ", ".join(summary)

    # ---------------- README links ----------------
    @check("local links in README.md resolve")
    def _links():
        text = (ROOT / "README.md").read_text(encoding="utf-8")
        targets = re.findall(r"\]\(([^)]+)\)", text)
        broken = []
        for target in targets:
            if target.startswith(("http://", "https://", "#", "mailto:")):
                continue
            path = ROOT / target.split("#")[0]
            if not path.exists():
                broken.append(target)
        assert not broken, f"broken local links: {broken}"
        return f"{len(targets)} links checked"

    # ---------------- reports carry no unfilled placeholders ----------------
    @check("no report contains an unfilled placeholder")
    def _placeholders():
        from src.report import MISSING
        offenders = []
        for markdown in sorted(config.REPORTS_DIR.glob("*.md")) + [ROOT / "README.md"]:
            if markdown.exists() and MISSING in markdown.read_text(encoding="utf-8"):
                offenders.append(markdown.name)
        assert not offenders, f"placeholder text remains in: {offenders}"
        return "all reports fully populated"

    # ---------------- write results ----------------
    config.ensure_dirs()
    failures = [r for r in RESULTS if r["status"] in {"FAIL", "ERROR"}]
    payload = {
        "passed": sum(1 for r in RESULTS if r["status"] == "pass"),
        "skipped": sum(1 for r in RESULTS if r["status"] == "skip"),
        "failed": len(failures),
        "checks": RESULTS,
    }
    (metrics / "artifact_verification.json").write_text(
        json.dumps(payload, indent=2), encoding="utf-8")

    print(f"\n{payload['passed']} passed, {payload['skipped']} skipped, "
          f"{payload['failed']} failed")
    print("wrote reports/metrics/artifact_verification.json")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
