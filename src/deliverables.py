"""
Builds the submission bundle in the formats the assignment accepts.

Produces, in ``deliverables/``:

* ``<Name>_AI_Capstone_Network_Intrusion_Detection_Report.docx``
  - every written report collated into one Word document, with a title page,
    a table-of-contents field, styled headings and tables, and each figure
    embedded at its first mention.
* ``<Name>_AI_Capstone_Technical_Presentation.pptx``
* ``<Name>_AI_Capstone_Executive_Presentation.pptx``
  - real PowerPoint decks built from `presentations/*.md`, with the recommended
    figure placed on the slide and the speaker notes in the notes pane.

Everything is derived from the generated Markdown, which is itself derived from
`reports/metrics/`. There is no third copy of any number.

Usage
-----
    python -m src.deliverables
    python -m src.deliverables --author "Arne Ramos"
"""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from src import config

DELIVERABLES_DIR = config.PROJECT_ROOT / "deliverables"
PRESENTATIONS_DIR = config.PROJECT_ROOT / "presentations"

DEFAULT_AUTHOR = "Arne Ramos"
PROJECT_TITLE = "Machine Learning-Based Network Intrusion Detection and Anomaly Classification"

#: Reports collated into the Word document, in reading order.
REPORT_ORDER: list[tuple[str, str]] = [
    ("problem_statement.md", "Part I — Problem Understanding and Framing"),
    ("dataset_documentation.md", "Part II — Data Collection and Understanding"),
    ("EDA_Feature_Engineering_Report.md", "Part III — EDA and Feature Engineering"),
    ("Model_Evaluation_Report.md", "Part IV — Model Implementation and Comparison"),
    ("Bias_Fairness_Analysis.md", "Part V — Critical Thinking, Ethical AI and Bias Audit"),
    ("Final_Project_Report.md", "Part VI — Final Project Report"),
    ("Generative_AI_Usage.md", "Part VII — Generative AI Usage Disclosure"),
    ("Rubric_Audit.md", "Part VIII — Rubric Audit"),
]

# Colour scheme, matching the figures (Okabe-Ito derived).
NAVY = (0x1F, 0x3A, 0x5F)
BLUE = (0x00, 0x72, 0xB2)
VERMILLION = (0xD5, 0x5E, 0x00)
GREY = (0x55, 0x55, 0x55)
LIGHT = (0xF2, 0xF2, 0xF2)


def _slug(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_")


# =========================================================================== #
# Minimal Markdown renderer for python-docx
# =========================================================================== #
INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))")
FIGURE_PATTERN = re.compile(r"(fig\d{2}[A-Za-z0-9_]*)\.png")


@dataclass
class DocxRenderer:
    """Renders a supported subset of Markdown into a python-docx document."""

    document: object
    used_figures: set = None

    def __post_init__(self):
        if self.used_figures is None:
            self.used_figures = set()

    # ---------------- inline ----------------
    def add_runs(self, paragraph, text: str, *, bold: bool = False,
                 italic: bool = False) -> None:
        """Apply bold / italic / inline-code / link-text formatting.

        Recurses into the interior of each span so that nested markup renders -
        the reports contain constructions like ``**`sttl` is a proxy**``, and a
        single non-recursive pass would emit the backticks literally.
        """
        from docx.shared import RGBColor

        for piece in INLINE_PATTERN.split(text):
            if not piece:
                continue
            if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
                self.add_runs(paragraph, piece[2:-2], bold=True, italic=italic)
            elif (piece.startswith("*") and piece.endswith("*")
                  and len(piece) > 2 and not piece.startswith("**")):
                self.add_runs(paragraph, piece[1:-1], bold=bold, italic=True)
            elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
                run = paragraph.add_run(piece[1:-1])
                run.font.name = "Consolas"
                run.font.color.rgb = RGBColor(*VERMILLION)
                run.bold, run.italic = bold, italic
            elif piece.startswith("[") and "](" in piece:
                label = piece[1:piece.index("](")]
                run = paragraph.add_run(label)
                run.font.color.rgb = RGBColor(*BLUE)
                run.underline = True
                run.bold, run.italic = bold, italic
            else:
                run = paragraph.add_run(piece)
                run.bold, run.italic = bold, italic

    # ---------------- blocks ----------------
    def heading(self, text: str, level: int) -> None:
        from docx.shared import Pt, RGBColor

        heading = self.document.add_heading(level=min(level, 4))
        heading.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
        heading.paragraph_format.space_after = Pt(6)
        self.add_runs(heading, text)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(*(NAVY if level <= 2 else BLUE))
            run.font.name = "Calibri"

    def paragraph(self, text: str) -> None:
        from docx.shared import Pt

        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(7)
        self.add_runs(paragraph, text)

    def bullet(self, text: str, level: int = 0) -> None:
        from docx.shared import Pt

        style = "List Bullet" if level == 0 else f"List Bullet {min(level + 1, 3)}"
        try:
            paragraph = self.document.add_paragraph(style=style)
        except KeyError:
            paragraph = self.document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        self.add_runs(paragraph, text)

    def numbered(self, text: str) -> None:
        from docx.shared import Pt

        paragraph = self.document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        self.add_runs(paragraph, text)

    def quote(self, text: str) -> None:
        from docx.shared import Pt, RGBColor

        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(24)
        paragraph.paragraph_format.space_after = Pt(8)
        self.add_runs(paragraph, text)
        for run in paragraph.runs:
            run.italic = True
            run.font.color.rgb = RGBColor(*GREY)

    def code_block(self, lines: list[str]) -> None:
        from docx.shared import Pt, RGBColor

        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.space_after = Pt(9)
        run = paragraph.add_run("\n".join(lines))
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    def table(self, rows: list[list[str]]) -> None:
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.shared import Pt, RGBColor

        if not rows:
            return
        width = max(len(r) for r in rows)
        table = self.document.add_table(rows=len(rows), cols=width)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r, row in enumerate(rows):
            for c in range(width):
                cell = table.cell(r, c)
                cell.text = ""
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(2)
                text = row[c] if c < len(row) else ""
                self.add_runs(paragraph, text)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
                    if r == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                if r == 0:
                    self._shade(cell, "1F3A5F")
                elif r % 2 == 0:
                    self._shade(cell, "F4F6F8")
        self.document.add_paragraph()

    @staticmethod
    def _shade(cell, hex_color: str) -> None:
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), hex_color)
        cell._tc.get_or_add_tcPr().append(shading)

    def figure(self, stem: str) -> None:
        """Embed a figure once, the first time it is referenced."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor

        if stem in self.used_figures:
            return
        path = config.FIGURES_DIR / f"{stem}.png"
        if not path.exists():
            return
        self.used_figures.add(stem)

        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(6.4))

        caption = self.document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption.add_run(f"Figure — {stem.replace('_', ' ')}")
        run.italic = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(*GREY)

    # ---------------- driver ----------------
    def render(self, markdown: str, base_heading: int = 1) -> None:
        lines = markdown.splitlines()
        index = 0
        while index < len(lines):
            line = lines[index]
            stripped = line.strip()

            if not stripped or set(stripped) <= {"-", "="} and len(stripped) >= 3:
                index += 1
                continue

            if stripped.startswith("```"):
                block, index = [], index + 1
                while index < len(lines) and not lines[index].strip().startswith("```"):
                    block.append(lines[index])
                    index += 1
                self.code_block(block)
                index += 1
                continue

            if stripped.startswith("|") and index + 1 < len(lines) and \
                    re.match(r"^\|[\s:|-]+\|$", lines[index + 1].strip()):
                rows = []
                header = [c.strip() for c in stripped.strip("|").split("|")]
                rows.append(header)
                index += 2
                while index < len(lines) and lines[index].strip().startswith("|"):
                    rows.append([c.strip() for c in lines[index].strip().strip("|").split("|")])
                    index += 1
                self.table(rows)
                continue

            match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if match:
                self.heading(match.group(2), len(match.group(1)) + base_heading - 1)
                index += 1
                continue

            if stripped.startswith(">"):
                quoted = []
                while index < len(lines) and lines[index].strip().startswith(">"):
                    quoted.append(lines[index].strip().lstrip(">").strip())
                    index += 1
                self.quote(" ".join(part for part in quoted if part))
                continue

            match = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
            if match:
                text, index = self._join_wrapped(lines, index + 1, match.group(2))
                self.bullet(text, level=len(match.group(1)) // 2)
                self._maybe_figure(text)
                continue

            match = re.match(r"^\s*\d+\.\s+(.*)$", line)
            if match:
                text, index = self._join_wrapped(lines, index + 1, match.group(1))
                self.numbered(text)
                self._maybe_figure(text)
                continue

            text, index = self._join_wrapped(lines, index + 1, stripped)
            self.paragraph(text)
            self._maybe_figure(text)

    @staticmethod
    def _starts_block(line: str) -> bool:
        """True if ``line`` opens a new Markdown block rather than continuing one."""
        stripped = line.strip()
        if not stripped:
            return True
        if set(stripped) <= {"-", "="} and len(stripped) >= 3:
            return True
        return bool(stripped.startswith(("```", "|", ">"))
                    or re.match(r"^#{1,6}\s", stripped)
                    or re.match(r"^[-*+]\s", stripped)
                    or re.match(r"^\d+\.\s", stripped))

    def _join_wrapped(self, lines: list[str], index: int, first: str) -> tuple[str, int]:
        """Collapse a source-wrapped block into one paragraph.

        The source Markdown is hard-wrapped at ~78 columns. Treating each
        physical line as its own paragraph produced ragged Word output and, more
        seriously, broke every inline span that straddled a line break - a
        ``**bold phrase`` opening on one line and closing on the next was
        emitted with its asterisks visible.
        """
        parts = [first]
        while index < len(lines) and not self._starts_block(lines[index]):
            parts.append(lines[index].strip())
            index += 1
        return " ".join(p for p in parts if p), index

    def _maybe_figure(self, text: str) -> None:
        for stem in FIGURE_PATTERN.findall(text):
            self.figure(stem)


# =========================================================================== #
# Word document
# =========================================================================== #
def build_docx(author: str = DEFAULT_AUTHOR) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Inches, Pt, RGBColor

    DELIVERABLES_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()

    # Page setup
    for section in document.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.95)
        section.right_margin = Inches(0.95)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    # ---------------- title page ----------------
    for _ in range(4):
        document.add_paragraph()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(PROJECT_TITLE)
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(*NAVY)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Graduate AI/ML Capstone Project  ·  Cybersecurity Domain")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(*VERMILLION)

    document.add_paragraph()
    rule = document.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = rule.add_run("—" * 34)
    run.font.color.rgb = RGBColor(*BLUE)

    for label, value in (
        ("Author", author),
        ("Date", date.today().strftime("%d %B %Y")),
        ("Dataset", "UNSW-NB15 (Moustafa & Slay, 2015)"),
        ("Task", "Supervised binary classification — benign vs. malicious network flow"),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"{label}:  ")
        run.bold = True
        run.font.size = Pt(11)
        run = paragraph.add_run(value)
        run.font.size = Pt(11)

    document.add_paragraph()
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(
        "Every metric in this document was produced by executed code and is read from\n"
        "reports/metrics/. No result, figure or statistic was written by hand.")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(*GREY)

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------- table of contents field ----------------
    document.add_heading("Table of Contents", level=1)
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    for tag, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, r' TOC \o "1-3" \h \z \u '),
        ("w:fldChar", {"w:fldCharType": "separate"}, None),
        ("w:t", {}, "Right-click here and choose “Update Field” to build the contents."),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        element = OxmlElement(tag)
        for key, value in attrs.items():
            element.set(qn(key), value)
        if text is not None:
            element.text = text
        run._r.append(element)

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------- body ----------------
    renderer = DocxRenderer(document)
    missing: list[str] = []

    for filename, part_title in REPORT_ORDER:
        path = config.REPORTS_DIR / filename
        if not path.exists():
            missing.append(filename)
            continue

        heading = document.add_heading(level=1)
        run = heading.add_run(part_title)
        run.font.color.rgb = RGBColor(*NAVY)

        markdown = path.read_text(encoding="utf-8")
        # The file's own H1 duplicates the part title we just wrote.
        markdown = re.sub(r"^#\s+.*?\n", "", markdown, count=1)
        renderer.render(markdown, base_heading=2)

        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------- figure appendix ----------------
    heading = document.add_heading(level=1)
    run = heading.add_run("Appendix A — Complete Figure Index")
    run.font.color.rgb = RGBColor(*NAVY)
    document.add_paragraph(
        "Every figure generated by the pipeline, in order. Figures already embedded "
        "in the body above are listed here for reference and are not repeated.")

    for path in sorted(config.FIGURES_DIR.glob("*.png")):
        stem = path.stem
        sub = document.add_heading(level=3)
        run = sub.add_run(stem.replace("_", " "))
        run.font.color.rgb = RGBColor(*BLUE)
        if stem in renderer.used_figures:
            paragraph = document.add_paragraph()
            run = paragraph.add_run("Embedded in the body above.")
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(*GREY)
        else:
            renderer.used_figures.discard(stem)
            renderer.figure(stem)

    # ---------------- footer ----------------
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"{author} — {PROJECT_TITLE}")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(*GREY)

    filename = f"{_slug(author)}_AI_Capstone_Network_Intrusion_Detection_Report.docx"
    path = DELIVERABLES_DIR / filename
    document.save(path)
    print(f"[deliverables] wrote deliverables/{filename}")
    if missing:
        print(f"[deliverables]   NOTE: these reports were not found and were skipped: "
              f"{', '.join(missing)}")
    return path


# =========================================================================== #
# PowerPoint decks
# =========================================================================== #
SLIDE_PATTERN = re.compile(r"^##\s+Slide\s+(\d+)\s*[—-]\s*(.*)$")


def parse_deck(markdown: str) -> list[dict]:
    """Parse a presentation Markdown file into slide dictionaries."""
    slides: list[dict] = []
    current: dict | None = None
    section: str | None = None

    for line in markdown.splitlines():
        stripped = line.strip()

        match = SLIDE_PATTERN.match(stripped)
        if match:
            if current:
                slides.append(current)
            current = {"number": int(match.group(1)), "title": match.group(2).strip(),
                       "content": [], "visual": "", "notes": [], "table": []}
            section = None
            continue

        if current is None:
            continue

        if stripped.startswith("**Content**"):
            section = "content"
            continue
        if stripped.startswith("**Recommended visual:**"):
            current["visual"] = stripped.split("**Recommended visual:**", 1)[1].strip()
            section = None
            continue
        if stripped.startswith("**Speaker notes:**"):
            current["notes"].append(
                stripped.split("**Speaker notes:**", 1)[1].strip())
            section = "notes"
            continue
        if stripped.startswith("---"):
            section = None
            continue
        if stripped.startswith("##"):
            section = None
            continue

        if section == "content":
            if stripped.startswith("|"):
                if not re.match(r"^\|[\s:|-]+\|$", stripped):
                    current["table"].append(
                        [c.strip() for c in stripped.strip("|").split("|")])
                continue
            bullet = re.match(r"^[-*+]\s+(.*)$", stripped)
            if bullet:
                # Only the single marker character - lstrip("-* ") would also eat
                # the opening ** of a bold lead-in, orphaning its closing pair.
                current["content"].append(bullet.group(1).strip())
            elif stripped:
                current["content"].append(stripped)
        elif section == "notes" and stripped:
            current["notes"].append(stripped)

    if current:
        slides.append(current)
    return slides


def _add_bullet_runs(paragraph, text: str, size, prefix: str = "") -> None:
    """Write ``text`` into a slide paragraph, honouring **bold** emphasis.

    Most bullets in the decks open with a bold lead-in ("**Objective:** ..."),
    which is the part a reader scans first; flattening it to plain text loses
    the visual hierarchy the slide was written around.
    """
    from pptx.dml.color import RGBColor

    body = RGBColor(0x2A, 0x2A, 0x2A)

    def emit(content: str, bold: bool) -> None:
        run = paragraph.add_run()
        run.text = content
        run.font.size = size
        run.font.bold = bold
        run.font.color.rgb = body

    if prefix:
        emit(prefix, False)
    for piece in (p for p in re.split(r"(\*\*.+?\*\*)", text) if p):
        bold = piece.startswith("**") and piece.endswith("**") and len(piece) > 4
        # _clean() strips edge whitespace, which would weld a bold lead-in to the
        # text following it; inside a run the surrounding spaces are significant.
        emit(_clean_inline(piece[2:-2] if bold else piece), bold)


def _clean_inline(text: str) -> str:
    """Strip Markdown emphasis, preserving leading/trailing whitespace."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return re.sub(r"\[([^\]]+?)\]\([^)]+?\)", r"\1", text)


def _clean(text: str) -> str:
    """Strip Markdown emphasis for slide text."""
    return _clean_inline(text).strip()


def _figure_for(visual: str) -> Path | None:
    match = FIGURE_PATTERN.search(visual)
    if not match:
        return None
    path = config.FIGURES_DIR / f"{match.group(1)}.png"
    return path if path.exists() else None


def build_pptx(source: Path, output_name: str, subtitle: str,
               accent: tuple[int, int, int], author: str) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    slides_data = parse_deck(source.read_text(encoding="utf-8"))
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)   # 16:9
    presentation.slide_height = Inches(7.5)

    blank = presentation.slide_layouts[6]
    width = presentation.slide_width
    height = presentation.slide_height

    def add_background_bar(slide, colour):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, width, Inches(0.09))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(*colour)
        bar.line.fill.background()
        bar.shadow.inherit = False

    def add_textbox(slide, left, top, box_width, box_height):
        box = slide.shapes.add_textbox(left, top, box_width, box_height)
        frame = box.text_frame
        frame.word_wrap = True
        return frame

    # ---------------- title slide ----------------
    first = slides_data[0] if slides_data else {"title": PROJECT_TITLE, "content": []}
    slide = presentation.slides.add_slide(blank)
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, width, Inches(2.9))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(*NAVY)
    banner.line.fill.background()
    banner.shadow.inherit = False

    frame = add_textbox(slide, Inches(0.9), Inches(0.85), width - Inches(1.8), Inches(1.9))
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = PROJECT_TITLE
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    paragraph = frame.add_paragraph()
    run = paragraph.add_run()
    run.text = subtitle
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor(*accent)

    frame = add_textbox(slide, Inches(0.9), Inches(3.3), width - Inches(1.8), Inches(3.0))
    for i, line in enumerate(first.get("content", [])[:5]):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = "•  " + _clean(line)
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        paragraph.space_after = Pt(9)

    frame = add_textbox(slide, Inches(0.9), height - Inches(1.0),
                        width - Inches(1.8), Inches(0.6))
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = f"{author}   ·   {date.today().strftime('%d %B %Y')}"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(*GREY)

    if first.get("notes"):
        slide.notes_slide.notes_text_frame.text = " ".join(
            _clean(n) for n in first["notes"])

    # ---------------- content slides ----------------
    for data in slides_data[1:]:
        slide = presentation.slides.add_slide(blank)
        add_background_bar(slide, accent)

        frame = add_textbox(slide, Inches(0.62), Inches(0.36),
                            width - Inches(1.24), Inches(0.95))
        paragraph = frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = _clean(data["title"])
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = RGBColor(*NAVY)

        figure = _figure_for(data.get("visual", ""))
        bullets = [b for b in data["content"] if b]
        table_rows = data.get("table") or []

        body_top = Inches(1.45)
        body_height = height - body_top - Inches(0.75)

        if figure is not None and bullets:
            text_width = Inches(6.1)
            image_left = Inches(6.95)
            image_width = width - image_left - Inches(0.62)
        elif figure is not None:
            text_width = Inches(0)
            image_left = Inches(1.6)
            image_width = width - Inches(3.2)
        else:
            text_width = width - Inches(1.24)
            image_left = image_width = None

        if bullets:
            frame = add_textbox(slide, Inches(0.62), body_top,
                                text_width if text_width else width - Inches(1.24),
                                body_height)
            size = Pt(15) if len(bullets) <= 5 else Pt(13)
            for i, bullet in enumerate(bullets):
                paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                _add_bullet_runs(paragraph, bullet, size, prefix="•  ")
                paragraph.space_after = Pt(8)

        if table_rows:
            top = body_top + (Inches(0.3) if not bullets else Inches(2.8))
            rows, cols = len(table_rows), max(len(r) for r in table_rows)
            shape = slide.shapes.add_table(
                rows, cols, Inches(0.62), top,
                (text_width if text_width else width - Inches(1.24)),
                Inches(0.32) * rows)
            table = shape.table
            for r, row in enumerate(table_rows):
                for c in range(cols):
                    cell = table.cell(r, c)
                    cell.text = _clean(row[c]) if c < len(row) else ""
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.alignment = PP_ALIGN.LEFT
                        for run in paragraph.runs:
                            run.font.size = Pt(10)
                            if r == 0:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        if figure is not None:
            from PIL import Image as PILImage

            with PILImage.open(figure) as image:
                aspect = image.height / image.width
            target_width = image_width
            target_height = Emu(int(target_width * aspect))
            available = height - body_top - Inches(0.7)
            if target_height > available:
                target_height = available
                target_width = Emu(int(target_height / aspect))
            left = image_left if bullets else int((width - target_width) / 2)
            slide.shapes.add_picture(str(figure), left,
                                     body_top + Inches(0.15),
                                     width=target_width, height=target_height)

        frame = add_textbox(slide, width - Inches(1.35), height - Inches(0.55),
                            Inches(0.9), Inches(0.35))
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.RIGHT
        run = paragraph.add_run()
        run.text = str(data["number"])
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(*GREY)

        notes = " ".join(_clean(n) for n in data.get("notes", []))
        if data.get("visual"):
            notes = f"[Visual: {_clean(data['visual'])}]\n\n{notes}"
        if notes.strip():
            slide.notes_slide.notes_text_frame.text = notes.strip()

    DELIVERABLES_DIR.mkdir(parents=True, exist_ok=True)
    path = DELIVERABLES_DIR / output_name
    presentation.save(path)
    print(f"[deliverables] wrote deliverables/{output_name} "
          f"({len(presentation.slides)} slides)")
    return path


# =========================================================================== #
def build_all(author: str = DEFAULT_AUTHOR) -> None:
    DELIVERABLES_DIR.mkdir(parents=True, exist_ok=True)
    slug = _slug(author)

    build_docx(author)

    decks = [
        ("technical_presentation_content.md",
         f"{slug}_AI_Capstone_Technical_Presentation.pptx",
         "Technical Deep-Dive  ·  ML / Data Science / Security Peers", BLUE),
        ("executive_presentation_content.md",
         f"{slug}_AI_Capstone_Executive_Presentation.pptx",
         "Executive Briefing  ·  Security Leadership", VERMILLION),
    ]
    for source_name, output_name, subtitle, accent in decks:
        source = PRESENTATIONS_DIR / source_name
        if not source.exists():
            print(f"[deliverables] SKIPPED {output_name}: "
                  f"presentations/{source_name} not found. "
                  "Run `python -m src.report` first.")
            continue
        build_pptx(source, output_name, subtitle, accent, author)

    _write_manifest(author)


def _write_manifest(author: str) -> None:
    slug = _slug(author)
    manifest = f"""# Submission Bundle

**{PROJECT_TITLE}**
**Author:** {author}   ·   **Generated:** {date.today().isoformat()}

---

## Files in this folder

| File | Format | What it is |
|---|---|---|
| `{slug}_AI_Capstone_Network_Intrusion_Detection_Report.docx` | Word | **All written responses collated into one document** — problem statement, dataset documentation, EDA, model evaluation, bias audit, final report, AI-usage disclosure and rubric audit, with every figure embedded. This is the file to upload if only one is accepted. |
| `{slug}_AI_Capstone_Technical_Presentation.pptx` | PowerPoint | 12-slide technical deck with figures on the slides and speaker notes in the notes pane. |
| `{slug}_AI_Capstone_Executive_Presentation.pptx` | PowerPoint | 10-slide executive deck — no equations, no code. |
| `DEMO_VIDEO_SCRIPT.md` | Markdown | Shot-by-shot script, timings, narration and recording setup for the demo video. |

## Still to do by hand

1. **Record the demo video** using `DEMO_VIDEO_SCRIPT.md`. Export as
   `{slug}_AI_Capstone_Demo_Video.mp4`.
2. **Open the Word document and update the Table of Contents** — right-click the
   contents field, choose *Update Field* → *Update entire table*. Word cannot
   populate it without being opened once.
3. **Check the decks in Presenter View** so you can see the speaker notes.
4. If the portal caps upload size, host the video (YouTube *Unlisted* or Drive)
   and paste the link onto slide 1 of the technical deck and the report cover page.

## Converting to PDF, if required

Open each file in Word or PowerPoint → **File → Export → Create PDF/XPS**.
Or from a terminal, if LibreOffice is installed:

```bash
soffice --headless --convert-to pdf deliverables/*.docx deliverables/*.pptx
```

## Regenerating this bundle

```bash
python -m src.pipeline --all     # recompute everything
python -m src.deliverables       # rebuild the .docx and .pptx files
```

Every number in these documents is read from `reports/metrics/`, which is
written by executed code — so the bundle cannot drift out of step with the
results.
"""
    path = DELIVERABLES_DIR / "README.md"
    path.write_text(manifest, encoding="utf-8")
    print("[deliverables] wrote deliverables/README.md")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the submission bundle.")
    parser.add_argument("--author", default=DEFAULT_AUTHOR)
    args = parser.parse_args(argv)
    build_all(author=args.author)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
